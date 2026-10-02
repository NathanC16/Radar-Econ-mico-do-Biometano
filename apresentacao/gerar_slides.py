"""Gera a apresentação oral do PI V-A (apresentacao/apresentacao_oral.pptx).

Uso: .venv/bin/python apresentacao/gerar_slides.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# paleta da marca 4WaTT (theme-4watt.css do site 4watt.tech)
TEAL = RGBColor(0x05, 0x7A, 0x64)    # verde 4WaTT, tom de texto (#03A589 escurecido p/ contraste)
DARK = RGBColor(0x2A, 0x07, 0x20)    # plum — títulos e fundo de capa
GRAY = RGBColor(0x6E, 0x64, 0x70)
LIGHT = RGBColor(0xF4, 0xF1, 0xEB)   # creme
AMBER = RGBColor(0x9A, 0x74, 0x00)   # dourado 4WaTT, tom de texto

W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def bar(s, color=RGBColor(0x03, 0xA5, 0x89), h=Inches(0.14)):
    r = s.shapes.add_shape(1, 0, 0, W, h)
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    return r


def title_box(s, text, sub=None):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.42), Inches(12.1), Inches(1.15))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = DARK
    if sub:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(15)
        r2.font.color.rgb = GRAY
    return tb


def body(s, lines, top=1.85, size=17, gap=6):
    """lines: list of (text, opts) or str."""
    tb = s.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12.0), H - Inches(top + 0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, opts = item
        else:
            text, opts = item, {}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(opts.get("gap", gap))
        r = p.add_run()
        r.text = text
        r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", False)
        r.font.color.rgb = opts.get("color", DARK)
    return tb


def card(s, x, y, w, h, title, value, note=None, color=TEAL):
    box = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = color
    box.line.width = Pt(1.2)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = value
    # 20pt: valores longos ("1.368 Mm³/d") cabem em 1 linha no card de 2.9"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = color
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = title
    r2.font.size = Pt(13)
    r2.font.color.rgb = DARK
    if note:
        p3 = tf.add_paragraph()
        r3 = p3.add_run()
        r3.text = note
        r3.font.size = Pt(11)
        r3.font.color.rgb = GRAY


def footer(s, n):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"PI V-A · Big Data e IA · PUC-Goiás — 4WaTT Bio Engenharia S/A        {n}"
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY


# ---------------------------------------------------------------- 1 capa
s = slide()
r = s.shapes.add_shape(1, 0, 0, W, H)
r.fill.solid()
r.fill.fore_color.rgb = DARK
r.line.fill.background()
bar(s, TEAL, Inches(0.2))
tb = s.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(3.4))
tf = tb.text_frame
p = tf.paragraphs[0]
run = p.add_run(); run.text = "Radar Econômico do Biometano"
run.font.size = Pt(44); run.font.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
p2 = tf.add_paragraph()
run2 = p2.add_run(); run2.text = "Painel analítico do mercado brasileiro de biogás/biometano e posicionamento da 4WaTT"
run2.font.size = Pt(20); run2.font.color.rgb = RGBColor(0x9E, 0xD9, 0xCC)
p3 = tf.add_paragraph(); p3.space_before = Pt(28)
run3 = p3.add_run(); run3.text = "Projeto Integrador V-A — Big Data e Inteligência Artificial"
run3.font.size = Pt(16); run3.font.color.rgb = RGBColor(0xC9, 0xB8, 0xC6)
p4 = tf.add_paragraph()
run4 = p4.add_run(); run4.text = "Nathan Crystiano França · outubro/2026"
run4.font.size = Pt(16); run4.font.color.rgb = RGBColor(0xC9, 0xB8, 0xC6)

# ---------------------------------------------------------------- 2 parceira
s = slide()
bar(s)
title_box(s, "A organização parceira", "4WaTT Bio Engenharia S/A — Goiânia/GO · bioengenharia de resíduos orgânicos")
body(s, [
    ("Converte resíduo em energia (biogás e biometano) com ciclo completo:", {"bold": True}),
    ("•  EVTE + business plan: estudo de viabilidade técnico-econômica e calculadora de pré-viabilidade", {}),
    ("•  EPC: construção turn-key de usinas de biodigestão e refinarias de biometano (BOT / EPCM)", {}),
    ("•  O&M: operação e manutenção recorrente com SCADA, automação e portal de indicadores", {}),
    ("•  Finance: assessoria de captação bancária para CAPEX", {}),
    ("Projetos públicos: CEASA Goiás (O&M) · Frigorífico Franca (2.800 Nm³/d em operação) · "
     "UTB Franca (biometano de curtume, em obra) · Organo Buritis", {"color": TEAL, "bold": True}),
    ("Nota de escopo: biometano = biogás refinado (~99% CH4, equivalente ao gás natural) — a camada regulada pela ANP onde o EPC/O&M é vendido; o core de biogás da 4WaTT segue outra regulação.",
     {"size": 11, "color": GRAY}),
])
footer(s, 2)

IMG = Path(__file__).resolve().parent / "img"
N = [2]  # contador de slides (capa e parceira já criados)


def novo(titulo, sub=None, cor=TEAL):
    s = slide()
    bar(s, cor)
    title_box(s, titulo, sub)
    N[0] += 1
    footer(s, N[0])
    return s


def tabela(s, linhas, top, larguras, size=13):
    """linhas[0] = cabeçalho."""
    shp = s.shapes.add_table(len(linhas), len(linhas[0]), Inches(0.7), Inches(top),
                             Inches(sum(larguras)), Inches(0.42 * len(linhas)))
    t = shp.table
    for j, w in enumerate(larguras):
        t.columns[j].width = Inches(w)
    for i, linha in enumerate(linhas):
        for j, txt in enumerate(linha):
            c = t.cell(i, j)
            c.text = txt
            c.fill.solid()
            c.fill.fore_color.rgb = DARK if i == 0 else (LIGHT if i % 2 else RGBColor(0xFF, 0xFF, 0xFF))
            for par in c.text_frame.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(size)
                    r.font.bold = i == 0
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else DARK
    return t


def imagem(s, arquivo, top=1.55, altura=5.35):
    pic = s.shapes.add_picture(str(IMG / arquivo), 0, Inches(top), height=Inches(altura))
    pic.left = int((W - pic.width) / 2)
    pic.line.color.rgb = RGBColor(0xD6, 0xCF, 0xC4)
    return pic


# ---------------------------------------------------------------- problema
s = novo("A necessidade identificada", "público: diretoria e equipes comercial e de engenharia da 4WaTT")
body(s, [
    ("O mercado de biometano vive expansão acelerada — e a 4WaTT precisa decidir, com dados:", {"bold": True}),
    ("1. Onde prospectar? Quais estados, municípios e substratos têm espaço para crescer?", {}),
    ("2. Que argumentos de mercado usar em EVTEs e apresentações a investidores e financiadores?", {}),
    ("3. O que acompanhar mês a mês para antecipar demanda por EPC e O&M?", {}),
    ("Hoje essa leitura está dispersa em boletins em PDF e notícias.", {"color": GRAY}),
    ("Solução: painel analítico open source com dados públicos, KPIs justificados, filtros e "
     "narrativa de data storytelling.", {"color": TEAL, "bold": True}),
])

# ---------------------------------------------------------------- interação
s = novo("Interação com a organização", "como o problema e a solução foram construídos com a 4WaTT")
body(s, [
    ("1.  Alinhamento do escopo com a equipe da 4WaTT: identificação do problema de inteligência de mercado", {}),
    ("2.  Definição conjunta dos KPIs e das fontes (boletim IEPUC/ANP, RenovaBio, site da empresa)", {}),
    ("3.  Validação intermediária do painel com a equipe", {}),
    ("4.  Apresentação interna dos resultados e coleta de feedback", {}),
    ("Datas, formatos e participantes: registrados no Relatório de Atividades Extensionistas e em "
     "documentos/evidencias/.", {"size": 14, "color": GRAY}),
    ("Vínculo formal: estágio curricular na 4WaTT + termo de formalização da atividade de extensão (PUC Goiás).",
     {"color": TEAL, "bold": True}),
], size=18, gap=10)

# ---------------------------------------------------------------- método
s = novo("Método e stack (100% open source)")
body(s, [
    ("Fontes: Boletim IEPUC-PUC-Rio ago/2026 (ANP · MME · B3) · dados abertos da ANP (produção, usinas e "
     "preços de combustíveis) · IBGE (rebanhos) · Panorama do Biometano em Goiás · site da 4WaTT", {}),
    ("", {}),
    ("Pandas → 16 bases (CSV)   |   Plotly → gráficos interativos   |   Streamlit → painel com 7 abas, filtros e tema escuro",
     {"bold": True}),
    ("Qualidade: 83 verificações automatizadas — busca de cada valor no texto dos PDFs e validação "
     "cruzada IEPUC × dados abertos da ANP (diferença < 0,5 mil m³/d em 2026)", {}),
    ("Reprodutível: requirements.txt + README + notebook de EDA + cópia das fontes em dados/brutos/", {}),
    ("Unidade do boletim: Mm³/d = mil m³ por dia (no painel: \"mil m³/d\").", {"size": 14, "color": GRAY}),
])

# ---------------------------------------------------------------- KPIs
s = novo("Indicadores (KPIs) e por que importam")
tabela(s, [
    ["KPI", "Decisão que apoia"],
    ["Produção nacional e crescimento (mil m³/d, % a.a.)", "Tamanho e ritmo da oportunidade"],
    ["Capacidade autorizada × em tramitação (ANP)", "Demanda futura por engenharia (EPC) e operação"],
    ["Concentração por estado e matéria-prima", "Onde o mercado ainda não está: regiões e substratos"],
    ["Utilização da capacidade (%)", "Usinas ociosas = demanda por O&M e otimização"],
    ["Contratos e transações anunciadas", "Prova de demanda para EVTE e investidores"],
    ["CBio: preço e intensidade de carbono", "Segunda receita por m³ na tese de viabilidade"],
    ["Potencial da pecuária por estado (IBGE)", "Goiás frente ao país: onde estão os dejetos"],
    ["Diesel × GNV por estado (ANP, R$/m³ eq.)", "Quanto o cliente economiza ao trocar de combustível"],
    ["Potencial de biogás em Goiás (fonte × município)", "Priorização da prospecção no estado-sede"],
], top=1.7, larguras=[6.0, 6.0], size=14)

# ---------------------------------------------------------------- mercado
s = novo("O mercado: produção em expansão acelerada",
         "média jan–ago/26 ~54% acima do mesmo período de 2025 (dados abertos ANP) · 7,6× desde jan/2020")
card(s, 0.7, 2.3, 2.9, 1.6, "Produção média jan–ago/26 (mil m³/d)", "466", "+54% vs jan–ago/25 (ANP)")
card(s, 3.8, 2.3, 2.9, 1.6, "Capacidade autorizada (mil m³/d)", "1.368", "697 em jun/25 → dobrou em 1 ano")
card(s, 6.9, 2.3, 2.9, 1.6, "Em tramitação — 49 pedidos (mil m³/d)", "2.023", "projeção IEPUC: 3.391 até dez/28")
card(s, 10.0, 2.3, 2.7, 1.6, "Utilização da capacidade", "38,7%", "sucroenergético: 27,2%", color=AMBER)
body(s, [
    ("Cada mil m³/d novo implica demanda por projeto, construção e operação — e a baixa utilização mostra que "
     "operar bem é um gargalo: o core da 4WaTT.", {"color": TEAL, "bold": True}),
], top=4.3)

# ---------------------------------------------------------------- concentração
s = novo("Concentração: onde o mercado já está")
body(s, [
    ("•  SP (33,6%) + RJ (26,3%) ≈ 60% da produção nacional; top 5 estados ≈ 99%", {"bold": True}),
    ("•  Substratos: aterros sanitários ≈ 80% · sucroenergético ≈ 19% · outros ≈ 1%", {}),
    ("•  Gás Verde = líder isolado: 139 mil m³/d ≈ 26% da produção nacional (ago/26)", {}),
    ("•  Só 9 estados têm usina autorizada; Norte nenhuma, Centro-Oeste uma (MS, sem produção)", {}),
    ("•  Goiás — sede da 4WaTT — NÃO tem nenhuma usina de biometano", {"color": AMBER, "bold": True}),
], size=19)
body(s, [
    ("Leitura: o crescimento veio concentrado em poucas UFs e num único substrato — o restante do mapa é fronteira.",
     {"color": TEAL, "bold": True}),
], top=4.9)

# ---------------------------------------------------------------- goiás
s = novo("Goiás: muito potencial, quase nenhuma produção",
         "Panorama do Biometano em Goiás — Governo de Goiás/CBIE (2026), potencial teórico")
card(s, 0.7, 2.2, 2.9, 1.6, "Potencial de biogás (bi m³/ano)", "2,7", "≈ 1,7 bi m³/ano de biometano")
card(s, 3.8, 2.2, 2.9, 1.6, "Sucroenergético", "72%", "vinhaça, torta de filtro, palha")
card(s, 6.9, 2.2, 2.9, 1.6, "Pecuária (mi m³/ano)", "402", "dispersa: top 30 municípios = 32%")
card(s, 10.0, 2.2, 2.7, 1.6, "Goiânia — RSU (mi m³/ano)", "124", "34% do RSU do estado", color=AMBER)
body(s, [
    ("•  Goiânia (sede e CEASA) é o maior potencial de RSU do estado → vitrine para replicar o modelo", {}),
    ("•  Pecuária dispersa → plantas descentralizadas de médio porte: formato EPC + O&M", {}),
    ("•  Sucroenergético (Goiatuba, Edéia, Mineiros, Caçu) → frente de maior escala, via parcerias", {}),
    ("•  Brasil (IBGE): GO é o 5º estado em potencial pecuário e o 2º em vacas ordenhadas", {}),
    ("•  Sem mercado de GNV em GO: 1 m³ de biometano substitui ≈ R$ 6,12 em diesel — acima do GNV de qualquer estado (set/26)", {}),
], top=4.2, size=17)

# ---------------------------------------------------------------- comercialização
s = novo("Comercialização e receita de carbono", "o mercado já compra — e o CBio soma uma segunda receita")
card(s, 0.7, 2.3, 2.9, 1.6, "Transações anunciadas (IEPUC)", "55", "GNC 71% · gasoduto 29%")
card(s, 3.8, 2.3, 2.9, 1.6, "Compradores industriais", "30 de 55", "setor industrial concentra a demanda")
card(s, 6.9, 2.3, 2.9, 1.6, "Preço médio CBio ago/26 (R$)", "24,53", "1 CBio = 1 tCO₂eq evitada")
card(s, 10.0, 2.3, 2.7, 1.6, "Usinas certificadas RenovaBio", "10", "19,5 mil CBios emitidos em ago/26")
body(s, [
    ("GNC (caminhão) viabiliza usinas longe de gasodutos — como no interior de Goiás. "
     "Argumento de EVTE: cada m³ vendido por usina certificada carrega dois fluxos — energia + carbono.",
     {"color": TEAL, "bold": True}),
], top=4.3)

# ---------------------------------------------------------------- portfolio
s = novo("O portfolio da 4WaTT no mapa")
body(s, [
    ("•  CEASA Goiás (Goiânia/GO): O&M de 600 t/mês de resíduos; R$ 1,6 mi investidos em novos projetos", {}),
    ("•  Frigorífico Franca (SP): projeto executivo de biodigestão — 2.800 Nm³/dia de biogás, em operação", {}),
    ("•  UTB Franca (SP): biometano a partir de resíduo de curtume — obras em estágio avançado, licença ambiental emitida", {}),
    ("•  Organo Buritis (Palmeiras de Goiás/GO): resíduo orgânico em biofertilizante e energia", {}),
], size=18)
body(s, [
    ("Cobre os três estágios do ciclo (operação → construção → projeto industrial em operação) e o mix de substratos "
     "minoritário no mercado nacional (~20% da produção fora de aterros).", {"color": TEAL, "bold": True}),
], top=4.9)

# ---------------------------------------------------------------- funcionamento
s = novo("Funcionamento do painel", "7 abas · tema claro/escuro · filtros na barra lateral e em cada aba · dados para download")
imagem(s, "painel_visao_geral.png")
s = novo("Funcionamento: Goiás em detalhe",
         "filtros por fonte de resíduo e quantidade de municípios; hover com valores")
imagem(s, "painel_goias.png")

# ---------------------------------------------------------------- visualização
s = novo("Escolhas de visualização")
tabela(s, [
    ["Pergunta", "Gráfico", "Por quê"],
    ["Como evolui a produção?", "Linha + referência 2025", "Tendência temporal; eixo a partir de zero"],
    ["Onde está a produção?", "Mapa coroplético", "Leitura territorial; sem dado ≠ valor baixo (bege)"],
    ["Qual a composição?", "Rosca (3 fatias)", "Partes de um todo com poucas categorias"],
    ["Quem produz mais?", "Barras horizontais ordenadas", "Rótulos longos; variação % no rótulo"],
    ["Onde está o potencial em GO?", "Barras empilhadas", "Ranking + composição por fonte"],
    ["Quanto vale o biometano?", "Linhas diesel × GNV", "Mesma base de comparação (R$/m³ eq.)"],
], top=1.7, larguras=[3.6, 3.6, 4.8], size=14)
body(s, [
    ("Paleta da marca 4WaTT com significado fixo: verde = sucroenergético/série principal, roxo = resíduo "
     "urbano, dourado = pecuária/outros · números no padrão brasileiro · "
     "títulos que afirmam o achado · blocos \"Insight para a 4WaTT\" em cada aba", {"size": 15, "color": GRAY}),
], top=5.2)

# ---------------------------------------------------------------- narrativa
s = novo("Narrativa (data storytelling)", "contexto → tensão → leitura → ação")
body(s, [
    ("1. Contexto — Visão Geral: o mercado cresce rápido (produção, capacidade, contratos)", {}),
    ("2. Tensão — Produção / Mercado: crescimento concentrado em 5 estados e em aterros; usinas ociosas", {}),
    ("3. Leitura — Goiás & Portfolio: o estado-sede tem potencial e quase nenhuma produção; "
     "o portfolio já atua nesse espaço", {}),
    ("4. Ação — blocos de insight: onde prospectar, como argumentar, o que monitorar", {}),
], size=19, gap=12)

# ---------------------------------------------------------------- insights
s = novo("Insights e recomendações", "o que o painel diz para a 4WaTT fazer", cor=AMBER)
body(s, [
    ("1. Priorizar Goiás: RSU em Goiânia (vitrine CEASA), pecuária dispersa (Rio Verde, Jataí) e "
     "parcerias sucroenergéticas no sul e sudoeste goianos", {"bold": True}),
    ("2. Oferecer O&M e otimização: o parque opera a 38,7% da capacidade autorizada", {}),
    ("3. Usar os KPIs em EVTEs e com investidores: crescimento + CBio = tese de viabilidade robusta", {}),
    ("4. Monitorar mensalmente o boletim IEPUC — o pipeline ANP é o termômetro da demanda por EPC/O&M", {}),
    ("Evoluções: pipeline comercial interno, preços de GLP/gás industrial/energia e atualização automática", {"size": 15, "color": GRAY}),
], size=18, gap=10)

# ---------------------------------------------------------------- extensão
s = novo("Contribuições e experiência extensionista")
body(s, [
    ("Para a 4WaTT", {"bold": True, "color": TEAL}),
    ("•  Painel reutilizável, bases validadas e recomendações acionáveis para prospecção e EVTE", {}),
    ("•  Rotina de monitoramento mensal do mercado com fontes públicas e rastreáveis", {}),
    ("Para a formação", {"bold": True, "color": TEAL}),
    ("•  BI e visualização aplicados a uma necessidade real, definida com a organização", {}),
    ("•  Trabalho com fontes heterogêneas (PDFs, sites), validação de dados e comunicação para decisores", {}),
    ("•  Relação universidade × empresa no setor de energia renovável e economia circular", {}),
], size=18, gap=8)

# ---------------------------------------------------------------- fim
s = slide()
r = s.shapes.add_shape(1, 0, 0, W, H)
r.fill.solid(); r.fill.fore_color.rgb = DARK; r.line.fill.background()
bar(s, TEAL, Inches(0.2))
tb = s.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(11.5), Inches(2.5))
p = tb.text_frame.paragraphs[0]
run = p.add_run(); run.text = "Obrigado."
run.font.size = Pt(44); run.font.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
p2 = tb.text_frame.add_paragraph()
run2 = p2.add_run()
run2.text = ("Código: github.com/NathanC16/Radar-Econ-mico-do-Biometano · Dados públicos: IEPUC-PUC-Rio, "
             "ANP, IBGE, Governo de Goiás e site da 4WaTT")
run2.font.size = Pt(15); run2.font.color.rgb = RGBColor(0xC9, 0xB8, 0xC6)

# metadados do arquivo (o modelo padrão da biblioteca traz valores genéricos)
cp = prs.core_properties
cp.title = "Radar Econômico do Biometano — Projeto Integrador V-A"
cp.author = cp.last_modified_by = "Nathan Crystiano França"
cp.subject = "Big Data e Inteligência Artificial — PUC Goiás · 4WaTT Bio Engenharia S/A"
cp.comments = cp.keywords = cp.category = ""
cp.revision = 1

out = Path(__file__).resolve().parent / "apresentacao_oral.pptx"
prs.save(out)
print("gerado:", out, "| slides:", len(prs.slides._sldIdLst))
